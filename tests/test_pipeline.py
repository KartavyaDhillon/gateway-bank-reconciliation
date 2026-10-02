"""End-to-end tests: synthetic data -> reconciliation -> API -> RAG analyst."""
import random
import sys
from datetime import date
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

import generate_data  # noqa: E402
import reconcile  # noqa: E402
from ai_analyst import ReconAnalyst, _amount_threshold  # noqa: E402


@pytest.fixture(scope="module")
def files(tmp_path_factory):
    d = tmp_path_factory.mktemp("data")
    gw, bank = generate_data.build(date(2026, 9, 30), 1250, random.Random(generate_data.SEED))
    gw.to_csv(d / "gw.csv", index=False)
    bank.to_csv(d / "bank.csv", index=False)
    return d / "gw.csv", d / "bank.csv", d


@pytest.fixture(scope="module")
def results(files):
    gw, bank, d = files
    res, summary, xlsx, _ = reconcile.run(gw, bank, d / "out", send=False)
    return res, summary, xlsx


def test_every_injected_problem_is_caught(results):
    _, summary, _ = results
    counts = dict(zip(summary.Status, summary.Count))
    assert counts == {"Matched": 1206, "Matched (T+1 credit)": 9, "Matched (amount + date)": 10,
                      "Amount Mismatch": 11, "Missing in Bank": 14, "Unknown Credit": 7,
                      "Duplicate Credit": 3}


def test_every_gateway_txn_accounted_for_once(results):
    res, _, _ = results
    ids = res["Txn ID"].dropna()
    assert len(ids) == 1250 and ids.is_unique


def test_excel_has_expected_sheets(results):
    from openpyxl import load_workbook
    wb = load_workbook(results[2], read_only=True)
    assert {"Summary", "Exceptions", "Amount Mismatch", "Missing in Bank", "All Results"} <= set(wb.sheetnames)


@pytest.mark.parametrize("q,expected", [("above 1 lakh", 1e5), ("over 50,000", 5e4), ("more than 2 cr", 2e7), ("hello", None)])
def test_amount_threshold_parser(q, expected):
    assert _amount_threshold(q) == expected


def test_rag_retrieves_large_missing_items_and_runbook(results):
    analyst = ReconAnalyst(results[0], ROOT / "knowledge" / "runbook.md")
    out = analyst.ask("Which missing in bank items are above 1 lakh and what should I do?")
    ids = [s["id"] for s in out["sources"]]
    assert "summary-missing-in-bank" in ids
    assert any(i.startswith("runbook") for i in ids)
    assert "208,322.01" in out["answer"] and "escalated" in out["answer"]


def test_llm_path_uses_retrieved_context(results, monkeypatch):
    analyst = ReconAnalyst(results[0], ROOT / "knowledge" / "runbook.md")
    captured = {}

    class FakeCompletions:
        def create(self, **kw):
            captured.update(kw)
            msg = type("M", (), {"content": "Escalate PG260930000143 [row-34]"})
            return type("R", (), {"choices": [type("C", (), {"message": msg})]})

    analyst.client = type("Client", (), {"chat": type("Chat", (), {"completions": FakeCompletions()})})
    out = analyst.ask("missing in bank above 1 lakh")
    assert out["mode"] == "llm"
    assert "Context:" in captured["messages"][1]["content"]
    assert captured["temperature"] == 0


def test_api_end_to_end(files, tmp_path, monkeypatch):
    import api
    monkeypatch.setattr(api, "RUNS_DIR", tmp_path)
    client = TestClient(api.app)
    gw, bank, _ = files
    r = client.post("/reconcile", files={"gateway": ("gw.csv", gw.read_bytes()), "bank": ("bank.csv", bank.read_bytes())})
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["gateway_rows"] == 1250 and body["exceptions"] == 35
    job = body["job_id"]
    assert client.get(f"/reconcile/{job}/report").status_code == 200
    exc = client.get(f"/reconcile/{job}/exceptions", params={"status": "Duplicate Credit"}).json()
    assert len(exc) == 3
    ans = client.post(f"/reconcile/{job}/ask", json={"question": "how many duplicate credits?"}).json()
    assert "3 items" in ans["answer"]


def test_api_rejects_non_csv(files):
    import api
    client = TestClient(api.app)
    r = client.post("/reconcile", files={"gateway": ("gw.pdf", b"x"), "bank": ("bank.csv", b"x")})
    assert r.status_code == 400
